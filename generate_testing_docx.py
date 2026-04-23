from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN_DARK = RGBColor(0x1b, 0x5e, 0x20)
GREEN_MID  = RGBColor(0x38, 0x8e, 0x3c)
GREEN_LIGHT= RGBColor(0xe8, 0xf5, 0xe9)
WHITE      = RGBColor(0xff, 0xff, 0xff)
MUTED      = RGBColor(0x61, 0x61, 0x61)

def set_cell_bg(cell, rgb):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

doc = Document()

# Margins
for section in doc.sections:
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# Default font
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)

# ── Chapter Title ──
h = doc.add_heading('', level=1)
run = h.add_run('Integration and Testing')
run.font.color.rgb = GREEN_DARK
run.font.size = Pt(16)
run.font.bold = True

# ── Intro paragraph ──
doc.add_paragraph(
    "Testing was carried out alongside the implementation of the app. Each module was tested "
    "independently before being combined and tested as a complete system."
)

# ── Types of testing heading ──
h2 = doc.add_heading('', level=2)
r2 = h2.add_run('Types of Testing Carried Out')
r2.font.color.rgb = GREEN_MID
r2.font.size = Pt(13)
r2.font.bold = True

# ── Bullet points ──
types = [
    ("Unit Testing",
     "Testing individual functions in isolation, such as the plant scoring formula and the calendar date calculation logic."),
    ("Integration Testing",
     "Testing how different modules work together, such as the Flutter app calling the FastAPI backend and receiving the correct response."),
    ("System Testing",
     "Testing the full app end-to-end as a complete working system across all features."),
    ("User Acceptance Testing",
     "Testing from the user's perspective to verify that all features behave as expected and that requirements have been met."),
]

for name, desc in types:
    p = doc.add_paragraph(style='List Bullet')
    r_bold = p.add_run(name + ' – ')
    r_bold.bold = True
    r_bold.font.color.rgb = GREEN_DARK
    p.add_run(desc)

doc.add_paragraph()

# ── Test Cases heading ──
h3 = doc.add_heading('', level=2)
r3 = h3.add_run('Test Cases')
r3.font.color.rgb = GREEN_MID
r3.font.size = Pt(13)
r3.font.bold = True

# ── Table ──
headers = ['Test Case ID', 'Test Case Name', 'Test Data', 'Expected Result', 'Actual Result']
col_widths = [Cm(2.0), Cm(4.2), Cm(4.5), Cm(4.8), Cm(2.5)]

rows = [
    ('TC_001', 'Login with valid credentials',
     'Email: ayush@gmail.com\nPassword: Yush123862868',
     'User logged in and redirected to home screen', 'As Expected'),

    ('TC_002', 'Login with invalid password',
     'Email: ayush@gmail.com\nPassword: wrongpass99',
     'Error message shown, user stays on login screen', 'As Expected'),

    ('TC_003', 'Login with empty fields',
     'Email: (empty)\nPassword: (empty)',
     'Validation error shown, login not attempted', 'As Expected'),

    ('TC_004', 'Register a new account',
     'Email: newuser@gmail.com\nPassword: Test@1234\nName: Ayush',
     'Account created, user redirected to home screen', 'As Expected'),

    ('TC_005', 'Logout',
     'Authenticated user session active',
     'User logged out and redirected to login screen', 'As Expected'),

    ('TC_006', 'Unauthenticated screen access',
     'User not logged in, navigates to home screen',
     'Automatically redirected to login screen', 'As Expected'),

    ('TC_007', 'Edit user profile',
     'Experience: Beginner\nPreference: Vegetables\nRegion: North',
     'Profile updated and saved successfully', 'As Expected'),

    ('TC_008', 'Browse plant catalog',
     'No filter applied',
     'Full list of Mauritius plants displayed', 'As Expected'),

    ('TC_009', 'Filter plants by type',
     'Filter: Vegetables',
     'Only vegetable plants shown in the catalog', 'As Expected'),

    ('TC_010', 'Global search — valid query',
     'Search: "tomato"',
     'Results returned from Perenual and displayed', 'As Expected'),

    ('TC_011', 'Global search — no match',
     'Search: "xyzplant123"',
     '"No global results" message displayed', 'As Expected'),

    ('TC_012', 'Upload garden photo for analysis',
     'Garden photo selected from gallery',
     'Zones detected and displayed on result screen', 'As Expected'),

    ('TC_013', 'Draw garden bed manually',
     'Name: Main Bed\nWidth: 200cm, Height: 150cm\nSun: Full Sun',
     'Bed created and saved, navigates to recommendations', 'As Expected'),

    ('TC_014', 'Get plant recommendations',
     'Bed: Full Sun\nRegion: North\nSeason: Summer',
     'Ranked list of recommended plants displayed', 'As Expected'),

    ('TC_015', 'Recommendations — offline AI',
     'Gemini and Ollama unavailable',
     'Rule-based recommendations displayed as fallback', 'As Expected'),

    ('TC_016', 'Generate garden layout',
     'Selected plants: Tomato, Basil, Lettuce',
     'Grid layout generated with plants placed in the bed', 'As Expected'),

    ('TC_017', 'Save garden layout',
     'Completed layout with placements',
     'Layout saved to Firestore, calendar triggered in background', 'As Expected'),

    ('TC_018', 'Companion planting — compatible',
     'Tomato placed next to Basil',
     'No warning shown, layout accepted', 'As Expected'),

    ('TC_019', 'Companion planting — incompatible',
     'Fennel placed next to Tomato',
     'Warning shown advising against the combination', 'As Expected'),

    ('TC_020', 'Generate planting calendar',
     'Layout saved with Tomato, Basil, Lettuce',
     'Sow, transplant, and harvest events created with correct dates', 'As Expected'),

    ('TC_021', 'Mark calendar event as complete',
     'Tap sow event for Tomato',
     'Event marked as done and UI updated', 'As Expected'),
]

table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'

# Set column widths and header row
hdr_row = table.rows[0]
for i, (header, width) in enumerate(zip(headers, col_widths)):
    cell = hdr_row.cells[i]
    cell.width = width
    cell.text = header
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = WHITE
    set_cell_bg(cell, GREEN_DARK)

# Data rows
for idx, row_data in enumerate(rows):
    row = table.add_row()
    bg = GREEN_LIGHT if idx % 2 == 1 else RGBColor(0xff, 0xff, 0xff)
    for i, (text, width) in enumerate(zip(row_data, col_widths)):
        cell = row.cells[i]
        cell.width = width
        cell.text = text
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(8.5)
                if i == 0:  # TC ID column bold
                    run.font.bold = True
                    run.font.color.rgb = GREEN_MID
        set_cell_bg(cell, bg)

doc.save(r'C:\Users\ayush\gardnx\testing_chapter.docx')
print('testing_chapter.docx saved.')
